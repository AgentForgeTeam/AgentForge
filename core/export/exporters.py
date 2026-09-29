"""Рендеринг результата в конкретные форматы (этап 8).

Все экспортёры получают одну и ту же модель блоков из ``bundle.py``,
поэтому содержание Markdown, DOCX и PDF совпадает по построению.

Отдельная история — кириллица в PDF: встроенные шрифты reportlab её не
знают, поэтому приходится искать в системе TrueType-шрифт. Если не нашли,
экспорт не молчит и не выдаёт кракозябры, а честно говорит об этом.
"""

from __future__ import annotations

import json
import logging
import re
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path

from core.export.bundle import Block, ExportOptions, ResultBundle, build_document

log = logging.getLogger("aiorc.export")


class ExportError(RuntimeError):
    """Ошибка экспорта с текстом, который можно показать пользователю."""


@dataclass
class ExportResult:
    path: Path
    format: str
    size: int
    note: str = ""


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def render_markdown(blocks: list[Block]) -> str:
    """Собирает Markdown-текст из блоков."""
    out: list[str] = []
    for block in blocks:
        if block.kind == "heading":
            out.append(f"{'#' * min(block.level, 6)} {block.text}")
        elif block.kind == "meta":
            out.append(f"*{block.text}*")
        elif block.kind == "text":
            out.append(block.text)
        elif block.kind == "code":
            out.append(f"```{block.language}\n{block.text}\n```")
        elif block.kind == "bullets":
            out.extend(f"- {item}" for item in block.items)
        elif block.kind == "divider":
            out.append("---")
        out.append("")
    return "\n".join(out).strip() + "\n"


def export_markdown(bundle: ResultBundle, options: ExportOptions,
                    path: Path) -> ExportResult:
    blocks = build_document(bundle, options)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_markdown(blocks), "utf-8")
    return ExportResult(path, "markdown", path.stat().st_size)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def export_docx(bundle: ResultBundle, options: ExportOptions,
                path: Path) -> ExportResult:
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Pt, RGBColor
    except ImportError as exc:
        raise ExportError(
            "Для экспорта в DOCX нужен пакет python-docx:\n"
            "pip install python-docx"
        ) from exc

    blocks = [_clean_block(b) for b in build_document(bundle, options)]
    document = Document()

    # Моноширинный стиль для кода — в стандартном шаблоне его нет.
    styles = document.styles
    try:
        code_style = styles.add_style("AiorcCode", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(9)
    except Exception:  # noqa: BLE001 — стиль уже есть
        code_style = styles["AiorcCode"]

    for block in blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=min(block.level, 4))
        elif block.kind == "meta":
            paragraph = document.add_paragraph(block.text)
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run("")
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x6D, 0x7A)
        elif block.kind == "text":
            for chunk in block.text.split("\n\n"):
                if chunk.strip():
                    document.add_paragraph(chunk.strip())
        elif block.kind == "code":
            paragraph = document.add_paragraph(block.text, style=code_style)
            paragraph.paragraph_format.left_indent = Pt(12)
        elif block.kind == "bullets":
            for item in block.items:
                document.add_paragraph(item, style="List Bullet")
        elif block.kind == "divider":
            document.add_paragraph("―" * 40)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return ExportResult(path, "docx", path.stat().st_size)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

#: где искать TrueType-шрифт с кириллицей
FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "C:/Windows/Fonts/segoeui.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
]
MONO_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf",
    "C:/Windows/Fonts/consola.ttf",
    "/System/Library/Fonts/Menlo.ttc",
]


def find_font(candidates: list[str]) -> Path | None:
    """Ищет шрифт в системе, а если не нашёл — в пакете matplotlib.

    matplotlib кладёт рядом с собой DejaVu, и это частый способ получить
    кириллический шрифт на машине, где системных TTF нет.
    """
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return path
    try:
        import matplotlib

        fonts_dir = Path(matplotlib.__file__).parent / "mpl-data" / "fonts" / "ttf"
        mono = "Mono" in " ".join(candidates)
        pattern = "DejaVuSansMono.ttf" if mono else "DejaVuSans.ttf"
        found = fonts_dir / pattern
        if found.exists():
            return found
    except Exception:  # noqa: BLE001
        pass
    return None


def export_pdf(bundle: ResultBundle, options: ExportOptions,
               path: Path) -> ExportResult:
    try:
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            HRFlowable,
            ListFlowable,
            ListItem,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )
    except ImportError as exc:
        raise ExportError(
            "Для экспорта в PDF нужен пакет reportlab:\n"
            "pip install reportlab"
        ) from exc

    regular = find_font(FONT_CANDIDATES)
    if regular is None:
        raise ExportError(
            "Не найден шрифт с поддержкой кириллицы, а встроенные шрифты PDF "
            "её не знают — текст получился бы нечитаемым.\n\n"
            "Установите шрифты DejaVu (Linux: пакет fonts-dejavu) либо "
            "выберите экспорт в DOCX или Markdown."
        )
    mono = find_font(MONO_CANDIDATES) or regular

    pdfmetrics.registerFont(TTFont("AiorcSans", str(regular)))
    pdfmetrics.registerFont(TTFont("AiorcMono", str(mono)))

    sheet = getSampleStyleSheet()
    body = ParagraphStyle("AiorcBody", parent=sheet["BodyText"],
                          fontName="AiorcSans", fontSize=10, leading=14,
                          alignment=TA_LEFT, spaceAfter=6)
    meta = ParagraphStyle("AiorcMeta", parent=body, fontSize=8,
                          textColor="#6a7080", spaceAfter=4)
    code_style = ParagraphStyle("AiorcCode", parent=body, fontName="AiorcMono",
                                fontSize=8, leading=11, leftIndent=8,
                                backColor="#f2f3f6", borderPadding=4)
    headings = {
        level: ParagraphStyle(
            f"AiorcH{level}", parent=body, fontName="AiorcSans",
            fontSize={1: 18, 2: 14, 3: 12, 4: 11}.get(level, 11),
            leading={1: 22, 2: 18, 3: 15, 4: 14}.get(level, 14),
            spaceBefore={1: 0, 2: 12, 3: 10, 4: 8}.get(level, 8), spaceAfter=6,
        )
        for level in (1, 2, 3, 4)
    }

    story: list = []
    for block in (_clean_block(b) for b in build_document(bundle, options)):
        if block.kind == "heading":
            story.append(Paragraph(_escape(block.text),
                                   headings.get(min(block.level, 4), body)))
        elif block.kind == "meta":
            story.append(Paragraph(_escape(block.text), meta))
        elif block.kind == "text":
            for chunk in block.text.split("\n\n"):
                if chunk.strip():
                    story.append(Paragraph(_escape(chunk.strip()).replace("\n", "<br/>"),
                                           body))
        elif block.kind == "code":
            story.append(Paragraph(
                _escape(block.text).replace("\n", "<br/>").replace(" ", "&nbsp;"),
                code_style,
            ))
            story.append(Spacer(1, 6))
        elif block.kind == "bullets" and block.items:
            story.append(ListFlowable(
                [ListItem(Paragraph(_escape(item), body)) for item in block.items],
                bulletType="bullet", start="•", leftIndent=14,
            ))
        elif block.kind == "divider":
            story.append(Spacer(1, 6))
            story.append(HRFlowable(width="100%", color="#d7dae2"))
            story.append(Spacer(1, 6))

    path.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=20 * mm, rightMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
        title=(bundle.task.title if bundle.task else bundle.workspace.name),
    )
    document.build(story)
    note = "" if mono != regular else "Моноширинный шрифт не найден, код набран основным."
    return ExportResult(path, "pdf", path.stat().st_size, note)


def _escape(raw: str) -> str:
    """Экранирует спецсимволы разметки reportlab."""
    return (raw or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


#: управляющие символы, недопустимые в XML (DOCX их не принимает вовсе):
#: всё ниже пробела, кроме табуляции и переводов строки
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
#: ANSI-последовательности цвета из вывода терминала: «\x1b[31m»
_ANSI = re.compile(r"\x1b\[[0-9;?]*[ -/]*[@-~]")


def _clean(text: str) -> str:
    """Убирает то, что сломало бы DOCX/PDF: цвета терминала и управляющие символы.

    Агенты вставляют в результат вывод программ как есть, а python-docx
    на первом же таком символе бросает исключение и экспорт целиком падает.
    """
    return _CONTROL.sub("", _ANSI.sub("", text or ""))


def _clean_block(block: Block) -> Block:
    return Block(block.kind, text=_clean(block.text), level=block.level,
                 items=[_clean(i) for i in block.items], language=block.language)


# ---------------------------------------------------------------------------
# ZIP
# ---------------------------------------------------------------------------


def export_zip(bundle: ResultBundle, options: ExportOptions,
               path: Path) -> ExportResult:
    """Архив с файлами проекта, отчётом в Markdown и машиночитаемым манифестом."""
    blocks = build_document(bundle, options)
    path.parent.mkdir(parents=True, exist_ok=True)

    manifest = {
        "workspace": bundle.workspace.name,
        "task": bundle.task.title if bundle.task else "",
        "generated_at": _now(),
        "subtasks": [
            {"title": s.title, "status": s.status,
             "agent": bundle.agent_name(s.agent_id, options.anonymize),
             "rework_count": s.rework_count}
            for s in bundle.subtasks
        ],
        "incidents": [
            {"kind": i.kind, "severity": i.severity, "status": i.status,
             "description": i.description, "resolution": i.resolution}
            for i in bundle.incidents
        ],
        "usage": {"tokens": bundle.tokens, "cost_usd": round(bundle.cost, 6)},
        "files": [],
    }

    root = bundle.workspace_dir
    written = 0
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("RESULT.md", render_markdown(blocks))

        if options.include_files:
            for file_path in bundle.files:
                try:
                    relative = file_path.relative_to(root)
                except ValueError:
                    continue
                arcname = str(Path("files") / relative)
                try:
                    archive.write(file_path, arcname)
                except OSError as exc:
                    log.warning("Не удалось добавить %s: %s", file_path, exc)
                    continue
                manifest["files"].append(str(relative))
                written += 1

        archive.writestr("manifest.json",
                         json.dumps(manifest, ensure_ascii=False, indent=2))

    note = (f"В архив добавлено файлов: {written}" if written
            else "Файлов в рабочем каталоге не было — в архиве только отчёт.")
    return ExportResult(path, "zip", path.stat().st_size, note)


def _now() -> str:
    from datetime import datetime

    return datetime.now().isoformat(timespec="seconds")


# ---------------------------------------------------------------------------
# Единая точка входа
# ---------------------------------------------------------------------------

EXPORTERS = {
    "markdown": export_markdown,
    "docx": export_docx,
    "pdf": export_pdf,
    "zip": export_zip,
}
EXTENSIONS = {"markdown": ".md", "docx": ".docx", "pdf": ".pdf", "zip": ".zip"}


def export(bundle: ResultBundle, options: ExportOptions, fmt: str,
           path: Path) -> ExportResult:
    """Экспортирует результат в указанном формате."""
    exporter = EXPORTERS.get(fmt)
    if exporter is None:
        raise ExportError(f"Неизвестный формат экспорта: {fmt}")
    return exporter(bundle, options, path)


def suggest_filename(bundle: ResultBundle, fmt: str) -> str:
    """Предлагает имя файла: название задачи + дата."""
    from datetime import datetime

    base = (bundle.task.title if bundle.task and bundle.task.title
            else bundle.workspace.name) or "result"
    safe = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in base).strip()
    safe = "_".join(safe.split())[:60] or "result"
    return f"{safe}_{datetime.now().strftime('%Y%m%d_%H%M')}{EXTENSIONS.get(fmt, '.txt')}"


def open_folder(path: Path) -> bool:
    """Открывает каталог с файлом в проводнике ОС."""
    import subprocess
    import sys

    folder = str(path.parent if path.is_file() else path)
    try:
        if sys.platform == "win32":
            subprocess.Popen(["explorer", folder])
        elif sys.platform == "darwin":
            subprocess.Popen(["open", folder])
        else:
            if shutil.which("xdg-open") is None:
                return False
            subprocess.Popen(["xdg-open", folder])
        return True
    except Exception:  # noqa: BLE001
        return False
